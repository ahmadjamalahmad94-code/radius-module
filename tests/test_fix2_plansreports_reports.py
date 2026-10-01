"""Fix wave 2 (stream plansreports) — finance / reports regressions.

* Web «المركز المالي» + «التقارير المالية» read the same source as the API
  (ledger payments + posted revenue_records), per currency, local-day filter.
* Open debt = outstanding (amount − posted settlements), not the loan value.
* Sales / payments / activations / P&L / snapshots: a voided payment and its
  reversal are not counted as transactions; totals split by currency.
* Arabic column labels (API ``columns`` + web table + CSV header), card-sales
  CSV always has a header row.
* Report date filters: local-day inclusive on the web reports and the
  operational-reports API; bad dates → 422 / flash; ``limit=abc`` → 422.

Run this file alone (per-file isolation).
"""
from __future__ import annotations

import csv
import io
import os
from datetime import datetime, timedelta, timezone
from uuid import uuid4

import pytest

TOKEN = "fix2-reports-token"
AUTH = {"Authorization": "Bearer " + TOKEN}
EXPIRE = datetime(2030, 1, 1, 12, 0, 0)


@pytest.fixture
def app(monkeypatch, tmp_path):
    db_file = os.path.join(tmp_path, "fix2_reports.db")
    monkeypatch.delenv("HOBERADIUS_ENV", raising=False)
    monkeypatch.delenv("FLASK_ENV", raising=False)
    monkeypatch.setenv("HOBERADIUS_DB_PATH", db_file)
    monkeypatch.setenv("HOBERADIUS_NO_WORKER", "1")
    monkeypatch.setenv("HOBERADIUS_NO_SEED", "1")
    monkeypatch.setenv("HOBERADIUS_LICENSE_GATE_TEST_BYPASS", "1")
    monkeypatch.setenv("HOBERADIUS_API_TOKENS", TOKEN)
    monkeypatch.setenv("FLASK_SECRET", "fix2-secret")
    monkeypatch.setenv("HOBERADIUS_BOOTSTRAP_ADMIN_USER", "owner_fix2")
    monkeypatch.setenv("HOBERADIUS_BOOTSTRAP_ADMIN_PASS", "owner-pass")
    from app.radius.db.connection import reset_for_tests
    reset_for_tests(db_file)
    from app import create_app
    application = create_app()
    with application.app_context():
        from app.radius.db.migrations_runner import run_pending_migrations
        from app.radius.db.repos import admins_repo, tenants_repo
        run_pending_migrations()
        tenants_repo.ensure_default_tenant()
        admins_repo.ensure_default_roles()
        yield application


@pytest.fixture
def client(app):
    return app.test_client()


# ─────────────── helpers ───────────────

def _web_login(client) -> str:
    res = client.post("/admin/radius/login",
                      data={"username": "owner_fix2", "password": "owner-pass"})
    assert res.status_code in {302, 303}, res.status_code
    client.get("/admin/radius/users")
    with client.session_transaction() as sess:
        return sess.get("_csrf_token", "")


def _db():
    from app.radius.db.connection import db
    return db()


def _plan(*, price=30.0, days=30) -> int:
    now = datetime.utcnow().isoformat()
    cur = _db().execute(
        "INSERT INTO access_plans(tenant_id, name, duration_minutes, validity_days, "
        "price, currency, enabled, created_at, updated_at) VALUES(1,?,?,?,?,?,1,?,?)",
        ("p_" + uuid4().hex[:6], days * 1440, days, price, "ILS", now, now))
    return int(cur.lastrowid)


def _sub(*, plan_id, username=None):
    from app.radius.core.types import Subscriber
    from app.radius.db.repos import subscribers_repo
    username = username or "u_" + uuid4().hex[:8]
    subscribers_repo.upsert_subscriber(Subscriber(
        id=None, tenant_id=1, username=username, password="secret", plan_id=plan_id,
        full_name="Fix2 User", mobile="0599000000", status="enabled", expire_at=EXPIRE))
    return subscribers_repo.get_subscriber(1, username)


def _data(res, status=200):
    body = res.get_json()
    assert res.status_code == status, body
    assert body["ok"] is True, body
    return body["data"]


def _err(res, status):
    body = res.get_json()
    assert res.status_code == status, body
    assert body["ok"] is False, body
    return body["error"]


def _pay(client, username, amount, currency="ILS") -> dict:
    return _data(client.post("/api/v1/payments", headers=AUTH, json={
        "username": username, "amount": amount, "currency": currency}), 201)["payment"]


def _ledger_payment(username, amount, currency, created_at) -> int:
    cur = _db().execute(
        "INSERT INTO accounting_ledger_entries(tenant_id, entry_type, direction, amount, "
        "currency, username, status, created_at) VALUES(1,'payment','credit',?,?,?,"
        "'posted',?)", (amount, currency, username, created_at))
    return int(cur.lastrowid)


def _revenue_record(collected_minor, profit_minor, currency="ILS", status="posted"):
    _db().execute(
        "INSERT INTO revenue_records(tenant_id, source_type, source_id, original_price_minor, "
        "retail_price_minor, wholesale_cost_minor, collected_amount_minor, net_profit_minor, "
        "company_share_minor, currency, status, metadata_json, created_at) "
        "VALUES(1,'card_user_purchase',1,?,?,0,?,?,?,?,?,'{}',?)",
        (collected_minor, collected_minor, collected_minor, profit_minor, profit_minor,
         currency, status, datetime.utcnow().isoformat() + "Z"))


def _local_today():
    from app.radius.core.system_config import local_today
    return local_today(1)


def _utc_of_local(day, hour, minute=0) -> str:
    """Local wall-clock (panel zone) → naive UTC ISO with Z."""
    from app.radius.core.system_config import tenant_tzinfo
    local = datetime(day.year, day.month, day.day, hour, minute, tzinfo=tenant_tzinfo(1))
    return local.astimezone(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ")


# ─────────────── item 5: finance center / financial report ───────────────

def test_finance_center_reads_ledger_payments_per_currency_api_and_web(client):
    from app.radius.services.business_os_finance_center import FinanceCenterService
    pid = _plan()
    s = _sub(plan_id=pid)
    _pay(client, s.username, 100)
    _pay(client, s.username, 40, "USD")
    voided = _pay(client, s.username, 7)
    _data(client.post(f"/api/v1/payments/{voided['id']}/void", headers=AUTH, json={}), 201)
    _revenue_record(1500, 500)  # a posted card-store sale: 15.00 collected, 5.00 profit

    summary = FinanceCenterService().dashboard(tenant_id=1)
    assert float(summary["total_collections"]) == 140.0  # 100 + 40 + 7 − 7
    assert float(summary["total_revenue"]) == 155.0      # + 15 card-store sale
    assert float(summary["total_profit"]) == 145.0       # payments + 5 profit
    by_cur = {c["currency"]: c for c in summary["revenue_by_currency"]}
    assert by_cur["ILS"]["revenue"] == 115.0 and by_cur["USD"]["revenue"] == 40.0
    assert summary["mixed_currency"] is True
    assert summary["payment_transactions"] == 2          # the voided one is not counted

    # The API reads the same source; a voided payment shows status voided and 0 profit.
    api = _data(client.get("/api/v1/finance/revenue", headers=AUTH))
    assert api["totals"]["revenue"] == 155.0 and api["totals"]["net_profit"] == 145.0
    row = next(i for i in api["items"] if i.get("metadata", {}).get("ledger_entry_id")
               and i["status"] == "voided")
    assert row["net_profit"] == 0.0 and row["company_share"] == 0.0

    _web_login(client)
    html = client.get("/admin/radius/finance-center").get_data(as_text=True)
    assert "40.00 USD" in html          # one figure per currency, not a mixed ₪ sum
    assert "115" in html
    rev_html = client.get("/admin/radius/finance-center?tab=revenue").get_data(as_text=True)
    assert "لا توجد سجلات إيراد بعد" not in rev_html
    assert s.username in rev_html


def test_finance_center_open_debt_is_the_outstanding(client):
    from app.radius.services.accounting import AccountingService
    from app.radius.services.business_os_finance_center import FinanceCenterService
    pid = _plan()
    s = _sub(plan_id=pid)
    loan = AccountingService(1).create_loan(
        {"username": s.username, "days": "2", "hours": "0", "currency": "ILS",
         "price_from_days": False, "amount": 10, "reason": "t", "apply_to_radius": False},
        actor="t")
    AccountingService(1).settle_loan(int(loan["id"]), {"amount": 4, "currency": "ILS"},
                                     actor="t")
    svc = FinanceCenterService()
    summary = svc.dashboard(tenant_id=1)
    debts = svc.debts(tenant_id=1)
    assert float(summary["total_loans"]) == 6.0 and float(summary["total_debts"]) == 6.0
    assert float(debts["total"]) == 6.0 and debts["items"][0]["outstanding"] == 6.0
    loans = svc.loans(tenant_id=1, status="open")
    assert loans[0]["outstanding"] == 6.0 and loans[0]["amount"] == 10.0

    _web_login(client)
    html = client.get("/admin/radius/finance-center?tab=loans_debts").get_data(as_text=True)
    assert "المتبقّي" in html and "6 ₪" in html
    users_html = client.get(f"/admin/radius/users/{s.username}/finance").get_data(as_text=True)
    assert "المتبقّي 6 ₪" in users_html  # per-subscriber open-loan chip = outstanding


def test_financial_report_uses_the_ledger_and_the_local_day(client):
    from app.radius.services.dashboard_reports import DashboardReportsService
    today = _local_today()
    # 01:30 local today is still «yesterday» in UTC — must count on the local day.
    _ledger_payment("night_owl", 7, "ILS", _utc_of_local(today, 1, 30))
    _ledger_payment("night_owl", 3, "USD", _utc_of_local(today, 23, 0))
    _ledger_payment("night_owl", 50, "ILS", _utc_of_local(today - timedelta(days=1), 12, 0))
    svc = DashboardReportsService(tenant_id=1)
    items = {i["metric"]: i for i in svc.report_data(
        "financial", date_from=today.isoformat(), date_to=today.isoformat())["items"]}
    assert items["payments"]["value"] == 10.0 and items["revenue"]["value"] == 10.0
    cur = {c["currency"]: c["total"] for c in items["payments"]["by_currency"]}
    assert cur == {"ILS": 7.0, "USD": 3.0}

    _web_login(client)
    html = client.get(f"/admin/radius/reports/financial?date_from={today}&date_to={today}"
                      ).get_data(as_text=True)
    assert "3.00 USD" in html and "130 ₪" not in html
    bad = client.get("/admin/radius/reports/financial?date_from=2026-13-45",
                     follow_redirects=True)
    assert bad.status_code == 200
    assert "صيغة التاريخ غير صحيحة" in bad.get_data(as_text=True)


# ─────────────── item 5 + 8: counts, currencies, labels ───────────────

def test_void_is_not_a_transaction_and_totals_split_by_currency(client):
    pid = _plan()
    s = _sub(plan_id=pid)
    _pay(client, s.username, 10)
    _pay(client, s.username, 5, "USD")
    voided = _pay(client, s.username, 3)
    _data(client.post(f"/api/v1/payments/{voided['id']}/void", headers=AUTH, json={}), 201)

    today = _local_today().isoformat()
    daily = _data(client.get("/api/v1/reports/sales/daily", headers=AUTH))
    row = next(r for r in daily["items"] if r["period"] == today)
    assert row["transactions"] == 2 and row["subscribers"] == 1
    assert row["total"] == 15.0 and row["mixed_currency"] is True
    assert {c["currency"]: c["total"] for c in row["by_currency"]} == {"ILS": 10.0, "USD": 5.0}
    labels = {c["key"]: c["label"] for c in daily["columns"]}
    assert labels["transactions"] == "عدد العمليات" and labels["avg_amount"] == "متوسّط العملية"

    pays = _data(client.get("/api/v1/reports/payments", headers=AUTH))
    mine = next(r for r in pays["items"] if r["username"] == s.username)
    assert mine["count"] == 2 and pays["totals"]["entries"] == 2
    acts = _data(client.get("/api/v1/reports/activations", headers=AUTH))
    assert next(r for r in acts["items"] if r["username"] == s.username)["activation_count"] == 2

    pl = _data(client.get("/api/v1/reports/profit-loss", headers=AUTH))["items"][0]
    pl_cur = {c["currency"]: c for c in pl["by_currency"]}
    assert pl_cur["ILS"]["net"] == 10.0 and pl_cur["USD"]["net"] == 5.0

    snap = _data(client.post("/api/v1/reports/snapshots", headers=AUTH,
                             json={"report_type": "daily"}), 201)["snapshot"]
    totals = {c["currency"]: c["total"] for c in snap["result"]["totals_by_currency"]}
    assert totals == {"ILS": 10.0, "USD": 5.0} and snap["result"]["mixed_currency"] is True


def test_web_accounting_reports_arabic_headers_and_per_currency(client):
    pid = _plan()
    s = _sub(plan_id=pid)
    _pay(client, s.username, 10)
    _pay(client, s.username, 5, "USD")
    _web_login(client)
    html = client.get("/admin/radius/finance/accounting?tab=reports&type=daily").get_data(as_text=True)
    assert "عدد العمليات" in html and "متوسّط العملية" in html and "عدد المشتركين" in html
    for raw in (">transactions<", ">avg_amount<", ">subscribers<"):
        assert raw not in html
    assert "5.00 USD" in html
    pl = client.get("/admin/radius/finance/accounting?tab=reports&type=profit_loss").get_data(as_text=True)
    assert "دفتر القيود المحاسبيّة" in pl and ">accounting_ledger_entries<" not in pl
    assert ">credits<" not in pl and "الصافي" in pl


def test_card_sales_csv_has_a_header_even_without_data(client):
    res = client.get("/api/v1/reports/card-sales/export.csv", headers=AUTH)
    assert res.status_code == 200
    text = res.get_data(as_text=True).lstrip("\ufeff")
    header = next(csv.reader(io.StringIO(text)))
    assert header == ["رقم الحزمة", "عدد العمليات", "الإجمالي"]
    _web_login(client)
    web = client.get("/admin/radius/finance/reports/export.csv?type=card_sales")
    assert web.status_code == 200 and "رقم الحزمة" in web.get_data(as_text=True)
    from openpyxl import load_workbook
    xlsx = client.get("/api/v1/reports/card-sales/export.xlsx", headers=AUTH)
    rows = list(load_workbook(io.BytesIO(xlsx.data), read_only=True).active.iter_rows(values_only=True))
    assert rows[0] == ("رقم الحزمة", "عدد العمليات", "الإجمالي")


# ─────────────── item 6: report date filters ───────────────

def _audit(action, target_type, created_at, actor="owner_fix2"):
    _db().execute(
        "INSERT INTO audit_log(tenant_id, actor, action, target_type, target_id, "
        "payload_json, created_at) VALUES(1,?,?,?,?,'{}',?)",
        (actor, action, target_type, "fix2_target", created_at))


def _cash_tx(username, amount, created_at):
    s = _sub(plan_id=_plan(), username=username)
    _db().execute(
        "INSERT INTO payment_transactions(tenant_id, subscriber_id, username, amount, currency, "
        "method, status, created_by, created_at) VALUES(1,?,?,?,'ILS','cash','posted','t',?)",
        (s.id, username, amount, created_at))


def test_web_report_date_filters_use_the_whole_local_day(client):
    today = _local_today()
    early = _utc_of_local(today, 1, 15)   # UTC «yesterday», local today
    late = _utc_of_local(today, 23, 30)   # after UTC midnight's cut, local today
    old = _utc_of_local(today - timedelta(days=2), 12, 0)
    _audit("update", "user", early)
    _audit("update", "user", late.replace("T", " ").replace("Z", ""))  # space format too
    _audit("update", "user", old)
    _cash_tx("fix2_cash_a", 11, early)
    _cash_tx("fix2_cash_b", 12, late)
    _cash_tx("fix2_cash_c", 13, old)
    _ledger_payment("fix2_bal", 21, "ILS", early)
    _web_login(client)
    d = today.isoformat()
    ue = client.get(f"/admin/radius/reports/user_events?date_from={d}&date_to={d}").get_data(as_text=True)
    assert ue.count("fix2_target") >= 2
    cash = client.get(f"/admin/radius/reports/cash_transactions?date_from={d}&date_to={d}"
                      ).get_data(as_text=True)
    assert "fix2_cash_a" in cash and "fix2_cash_b" in cash and "fix2_cash_c" not in cash
    bal = client.get(f"/admin/radius/reports/balance_movements?date_from={d}&date_to={d}"
                     ).get_data(as_text=True)
    assert "fix2_bal" in bal
    mgr = client.get(f"/admin/radius/reports/manager_events?date_from={d}&date_to={d}")
    assert mgr.status_code == 200
    only_to = client.get(f"/admin/radius/reports/cash_transactions?date_to={d}").get_data(as_text=True)
    assert "fix2_cash_b" in only_to  # `to` no longer cuts at UTC midnight
    bad = client.get("/admin/radius/reports/cash_transactions?date_from=garbage",
                     follow_redirects=True).get_data(as_text=True)
    assert "صيغة التاريخ غير صحيحة" in bad and "fix2_cash_c" in bad  # flash + unfiltered


def test_operational_reports_api_honours_dates_and_strict_limit(client):
    today = _local_today()
    _cash_tx("fix2_api_a", 11, _utc_of_local(today, 1, 15))
    _cash_tx("fix2_api_b", 12, _utc_of_local(today - timedelta(days=3), 12, 0))
    _ledger_payment("fix2_api_bal", 21, "ILS", _utc_of_local(today, 2, 0))
    d = today.isoformat()
    got = _data(client.get(f"/api/v1/operational-reports/cash-transactions?from={d}&to={d}",
                           headers=AUTH))
    names = {i["username"] for i in got["items"]}
    assert "fix2_api_a" in names and "fix2_api_b" not in names
    future = _data(client.get("/api/v1/operational-reports/balance-movements?date_from=2030-01-01",
                              headers=AUTH))
    assert future["count"] == 0
    today_bal = _data(client.get(
        f"/api/v1/operational-reports/balance-movements?date_from={d}&date_to={d}", headers=AUTH))
    assert any(i["username"] == "fix2_api_bal" for i in today_bal["items"])
    err = _err(client.get("/api/v1/operational-reports/cash-transactions?from=2026-02-31",
                          headers=AUTH), 422)
    assert "صيغة التاريخ" in err["message"]
    err = _err(client.get(f"/api/v1/operational-reports/cash-transactions?from={d}&to=2000-01-01",
                          headers=AUTH), 422)
    assert "يسبق" in err["message"]
    err = _err(client.get("/api/v1/operational-reports/cash-transactions?limit=abc",
                          headers=AUTH), 422)
    assert "limit" in err["message"] and "رقمًا" in err["message"]
