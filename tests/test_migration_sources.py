"""اختبارات طبقة المصادر في محرّك الترحيل — الفحص والاستخراج (نقيّة، بلا DB).

تغطّي: كشف النوع من المحتوى، فحص قاعدة SQLite، تحليل تفريغ SQL (MySQL/PG)،
قراءة CSV/Excel، وتصدير MikroTik ‎.rsc.

شغّل هذا الملف وحده (عزل الاختبارات لكل ملف).
"""
from __future__ import annotations

import io
import os
import sqlite3
import tempfile

from app.radius.services.migration import sources


# ── كشف النوع ────────────────────────────────────────────────────────

class TestSniff:
    def test_sqlite_magic(self):
        assert sources.sniff_source(b"SQLite format 3\x00rest", "x") == "sqlite"

    def test_pdf_magic(self):
        assert sources.sniff_source(b"%PDF-1.7\n...", "x") == "pdf"

    def test_xlsx_zip(self):
        body = b"PK\x03\x04" + b"......[Content_Types].xml...xl/workbook.xml"
        assert sources.sniff_source(body, "x") == "xlsx"

    def test_sql_dump_by_content(self):
        assert sources.sniff_source(b"-- dump\nINSERT INTO t VALUES (1);", "x") == "sql_dump"

    def test_mikrotik_by_content(self):
        assert sources.sniff_source(b"/ppp secret\nadd name=a password=b\n", "x") == "mikrotik"

    def test_csv_by_content(self):
        assert sources.sniff_source(b"a,b,c\n1,2,3\n", "x") == "csv"

    def test_extension_fallback(self):
        assert sources.sniff_source(b"weird\x01\x02bytes", "data.csv") == "csv"
        assert sources.sniff_source(b"weird\x01\x02bytes", "router.rsc") == "mikrotik"

    def test_unknown(self):
        assert sources.sniff_source(b"\x00\x01\x02\x03", "x") == "unknown"
        assert sources.sniff_source(b"", "x") == "unknown"


# ── CSV ──────────────────────────────────────────────────────────────

class TestCsv:
    def test_header_and_rows(self):
        ds = sources.introspect(b"username,password\nali,1\nsara,2\n", "u.csv")
        assert ds.fmt == "csv"
        t = ds.tables[0]
        assert t.columns == ["username", "password"]
        assert t.row_count == 2
        assert t.rows[0] == {"username": "ali", "password": "1"}

    def test_semicolon_delimiter(self):
        ds = sources.introspect(b"a;b;c\n1;2;3\n", "x.csv")
        assert ds.tables[0].columns == ["a", "b", "c"]

    def test_headerless_numeric(self):
        # كل الصفوف رقميّة → لا ترويسة، أعمدة مُولَّدة.
        ds = sources.introspect(b"1,2\n3,4\n", "x.csv")
        t = ds.tables[0]
        assert t.columns == ["col1", "col2"]
        assert t.row_count == 2

    def test_arabic_cp1256(self):
        body = "اسم,كلمة\nعلي,1\n".encode("cp1256")
        ds = sources.introspect(body, "ar.csv")
        assert ds.tables[0].row_count == 1
        assert "علي" in ds.tables[0].rows[0].values()


# ── SQLite ───────────────────────────────────────────────────────────

def _make_sqlite(script: str) -> bytes:
    fd, path = tempfile.mkstemp(suffix=".db")
    os.close(fd)
    try:
        conn = sqlite3.connect(path)
        conn.executescript(script)
        conn.commit()
        conn.close()
        with open(path, "rb") as fh:
            return fh.read()
    finally:
        os.unlink(path)


class TestSqlite:
    def test_introspect_tables_and_rows(self):
        data = _make_sqlite("""
            CREATE TABLE users (id INTEGER, name TEXT, pw TEXT);
            INSERT INTO users VALUES (1,'ali','x'),(2,'sara','y');
            CREATE TABLE plans (id INTEGER, title TEXT);
            INSERT INTO plans VALUES (1,'Gold');
        """)
        ds = sources.introspect(data, "db.sqlite")
        assert ds.fmt == "sqlite"
        names = {t.name for t in ds.tables}
        assert {"users", "plans"} <= names
        users = ds.table("users")
        assert users.columns == ["id", "name", "pw"]
        assert users.row_count == 2
        assert users.rows[0]["name"] == "ali"

    def test_corrupt_sqlite_warns(self):
        ds = sources.introspect(b"SQLite format 3\x00garbage-not-a-db", "x.db")
        assert ds.fmt == "sqlite"
        assert ds.warnings  # رسالة خطأ ودّية، بلا انهيار


# ── SQL dump ─────────────────────────────────────────────────────────

class TestSqlDump:
    def test_mysql_backticks(self):
        dump = b"""
        CREATE TABLE `subs` (`id` int, `username` varchar(50), `pass` varchar(50));
        INSERT INTO `subs` (`id`,`username`,`pass`) VALUES (1,'u1','p1'),(2,'u2','p2');
        """
        ds = sources.introspect(dump, "d.sql")
        assert ds.fmt == "sql_dump"
        t = ds.table("subs")
        assert t.columns == ["id", "username", "pass"]
        assert t.row_count == 2
        assert t.rows[1] == {"id": "2", "username": "u2", "pass": "p2"}

    def test_postgres_quotes_and_null(self):
        dump = b'''
        CREATE TABLE accounts (id integer, email text, note text);
        INSERT INTO accounts VALUES (1, 'a@b.com', NULL);
        '''
        ds = sources.introspect(dump, "pg.sql")
        t = ds.table("accounts")
        assert t.rows[0]["email"] == "a@b.com"
        assert t.rows[0]["note"] == ""           # NULL → فارغ

    def test_escaped_quote_in_value(self):
        dump = b"""CREATE TABLE t (id int, v varchar);
        INSERT INTO t VALUES (1, 'O''Brien'),(2, 'line');"""
        ds = sources.introspect(dump, "q.sql")
        t = ds.table("t")
        assert t.rows[0]["v"] == "O'Brien"

    def test_multiple_inserts_same_table(self):
        dump = b"""CREATE TABLE t (id int, u varchar);
        INSERT INTO t VALUES (1,'a');
        INSERT INTO t VALUES (2,'b'),(3,'c');"""
        ds = sources.introspect(dump, "m.sql")
        assert ds.table("t").row_count == 3


# ── MikroTik ─────────────────────────────────────────────────────────

class TestMikrotik:
    def test_ppp_and_hotspot(self):
        rsc = b"""# RouterOS export
/ppp secret
add name=ahmad password=pp1 profile=10M service=pppoe
add name=omar password=pp2 profile=5M
/ip hotspot user
add name=guest password=g1 profile=1h
"""
        ds = sources.introspect(rsc, "e.rsc")
        assert ds.fmt == "mikrotik"
        ppp = ds.table("ppp_secrets")
        assert ppp.row_count == 2
        assert ppp.rows[0]["name"] == "ahmad"
        assert ppp.rows[0]["profile"] == "10M"
        assert ds.table("hotspot_users").row_count == 1

    def test_quoted_values_and_continuation(self):
        rsc = b'/ppp secret\nadd name=a password="p p" comment="hi there"\n'
        ds = sources.introspect(rsc, "e.rsc")
        row = ds.table("ppp_secrets").rows[0]
        assert row["password"] == "p p"
        assert row["comment"] == "hi there"


# ── Excel ────────────────────────────────────────────────────────────

class TestXlsx:
    def test_sheets_as_tables(self):
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.title = "subs"
        ws.append(["username", "password"])
        ws.append(["ali", 1234])
        ws.append(["sara", 5678])
        buf = io.BytesIO()
        wb.save(buf)
        ds = sources.introspect(buf.getvalue(), "book.xlsx")
        assert ds.fmt == "xlsx"
        t = ds.table("subs")
        assert t.columns == ["username", "password"]
        assert t.rows[0] == {"username": "ali", "password": "1234"}  # 1234 لا 1234.0
