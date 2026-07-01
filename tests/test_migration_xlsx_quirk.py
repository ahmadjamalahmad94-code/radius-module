"""اختبارات صلابة قراءة Excel — بما فيها ملفّات بأنماط مشوّهة تُفشل openpyxl
(«expected Fill»)، وتتطلّب مسار XML الخام. + جُمل مشتركة + inline strings.

شغّل هذا الملف وحده."""
from __future__ import annotations

import io
import zipfile

from app.radius.services.migration import sources


def _openpyxl_book() -> bytes:
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "subs"
    ws.append(["username", "password", "plan"])
    ws.append(["ali", 1234, "Gold"])
    ws.append(["sara", 5678, "Silver"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _corrupt_styles(xlsx_bytes: bytes) -> bytes:
    """يُفسد xl/styles.xml (يحاكي «expected Fill») ويُعيد بناء الحزمة."""
    src = zipfile.ZipFile(io.BytesIO(xlsx_bytes))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for item in src.namelist():
            data = src.read(item)
            if item == "xl/styles.xml":
                data = b"<styleSheet><fills><fill/></fills></styleSheet>"  # ناقص/مشوّه
            z.writestr(item, data)
    return out.getvalue()


class TestXlsxRobust:
    def test_normal_xlsx(self):
        ds = sources.introspect(_openpyxl_book(), "b.xlsx")
        t = ds.table("subs")
        assert t.columns == ["username", "password", "plan"]
        assert t.rows[0] == {"username": "ali", "password": "1234", "plan": "Gold"}

    def test_corrupt_styles_falls_back_to_raw_xml(self):
        corrupted = _corrupt_styles(_openpyxl_book())
        ds = sources.introspect(corrupted, "b.xlsx")
        assert ds.fmt == "xlsx"
        t = ds.table("subs") or (ds.tables[0] if ds.tables else None)
        assert t is not None, ds.warnings
        assert t.row_count == 2
        # الأعمدة قُرئت رغم عطب الأنماط.
        assert "username" in t.columns

    def test_raw_xml_reader_directly(self):
        grids = sources._xlsx_via_raw_xml(_openpyxl_book())
        assert grids
        name, grid = grids[0]
        assert grid[0] == ["username", "password", "plan"]

    def test_col_ref_index(self):
        assert sources._col_ref_to_index("A1") == 0
        assert sources._col_ref_to_index("B2") == 1
        assert sources._col_ref_to_index("AA1") == 26
        assert sources._col_ref_to_index("AB3") == 27


class TestXlsxInlineStrings:
    def test_inline_strings(self):
        # ابنِ xlsx بسيطًا يدويًّا بقيَم inline (t="inlineStr").
        sheet = (
            '<?xml version="1.0"?>'
            '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
            '<sheetData>'
            '<row r="1"><c r="A1" t="inlineStr"><is><t>name</t></is></c>'
            '<c r="B1" t="inlineStr"><is><t>price</t></is></c></row>'
            '<row r="2"><c r="A2" t="inlineStr"><is><t>Gold</t></is></c>'
            '<c r="B2"><v>10</v></c></row>'
            '</sheetData></worksheet>')
        wb = ('<?xml version="1.0"?><workbook '
              'xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
              'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
              '<sheets><sheet name="S1" sheetId="1" r:id="rId1"/></sheets></workbook>')
        rels = ('<?xml version="1.0"?><Relationships '
                'xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
                '<Relationship Id="rId1" Target="worksheets/sheet1.xml"/></Relationships>')
        ct = ('<?xml version="1.0"?><Types '
              'xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>')
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, "w") as z:
            z.writestr("[Content_Types].xml", ct)
            z.writestr("xl/workbook.xml", wb)
            z.writestr("xl/_rels/workbook.xml.rels", rels)
            z.writestr("xl/worksheets/sheet1.xml", sheet)
        ds = sources.introspect(buf.getvalue(), "inline.xlsx")
        t = ds.tables[0]
        assert t.columns == ["name", "price"]
        assert t.rows[0] == {"name": "Gold", "price": "10"}
